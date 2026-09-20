"""Slide shapes, tables and available speaker notes without rendering."""

from pathlib import Path
from itertools import islice

from pptx import Presentation

from ordenia.analysis.models import AnalysisLimits, ExtractedContent


class PptxExtractor:
    def extract(self, path: Path, limits: AnalysisLimits) -> ExtractedContent:
        presentation = Presentation(path)
        parts: list[str] = []
        remaining = limits.max_extracted_characters
        truncated = len(presentation.slides) > limits.max_pptx_slides

        def add(value: str) -> None:
            nonlocal remaining, truncated
            if not value.strip():
                return
            if remaining <= 0:
                truncated = True
                return
            parts.append(value[:remaining])
            remaining -= len(parts[-1])
            truncated |= len(value) > len(parts[-1])

        for slide in islice(presentation.slides, limits.max_pptx_slides):
            for shape in slide.shapes:
                if shape.has_text_frame:
                    add(shape.text)
                if shape.has_table:
                    for row in shape.table.rows:
                        add(" | ".join(cell.text for cell in row.cells))
                if truncated:
                    break
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
                add(slide.notes_slide.notes_text_frame.text)
            if truncated:
                break
        props = presentation.core_properties
        return ExtractedContent("\n".join(parts), "PPTX", props.title or "", props.author or "",
                                props.subject or "", slide_count=len(presentation.slides), truncated=truncated)

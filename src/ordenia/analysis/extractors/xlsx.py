"""Stream workbook values with explicit cell and text limits."""

from pathlib import Path

from openpyxl import load_workbook

from ordenia.analysis.models import AnalysisLimits, ExtractedContent


class XlsxExtractor:
    def extract(self, path: Path, limits: AnalysisLimits) -> ExtractedContent:
        workbook = load_workbook(path, read_only=True, data_only=True)
        parts: list[str] = []
        remaining = limits.max_extracted_characters
        cells = 0
        truncated = False
        wide_columns = False
        try:
            for sheet in workbook.worksheets:
                parts.append(sheet.title[:remaining])
                remaining -= len(parts[-1])
                wide_columns |= bool(sheet.max_column and sheet.max_column > limits.max_xlsx_columns)
                for row in sheet.iter_rows(max_col=min(sheet.max_column or 1, limits.max_xlsx_columns), values_only=True):
                    if remaining <= 0 or cells >= limits.max_xlsx_cells:
                        truncated = True
                        break
                    values: list[str] = []
                    for value in row:
                        cells += 1
                        if cells > limits.max_xlsx_cells:
                            truncated = True
                            break
                        if value is not None:
                            values.append(str(value)[:remaining])
                    if values:
                        line = " | ".join(values)
                        parts.append(line[:remaining])
                        remaining -= len(parts[-1])
                        truncated |= len(line) > len(parts[-1])
                    if truncated:
                        break
                if truncated:
                    break
            props = workbook.properties
            return ExtractedContent("\n".join(parts), "XLSX", props.title or "", props.creator or "",
                                    props.subject or "", truncated=truncated or wide_columns)
        finally:
            workbook.close()

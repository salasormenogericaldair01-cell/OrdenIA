"""DOCX paragraphs, headers and table text without format preservation."""

from pathlib import Path

from docx import Document

from ordenia.analysis.models import AnalysisLimits, ExtractedContent


class DocxExtractor:
    def extract(self, path: Path, limits: AnalysisLimits) -> ExtractedContent:
        document = Document(path)
        parts: list[str] = []
        remaining = limits.max_extracted_characters
        truncated = False

        def add(text: str) -> None:
            nonlocal remaining, truncated
            if not text.strip():
                return
            if remaining <= 0:
                truncated = True
                return
            parts.append(text[:remaining])
            remaining -= len(parts[-1])
            truncated |= len(text) > len(parts[-1])

        for section in document.sections:
            for paragraph in section.header.paragraphs:
                add(paragraph.text)
        for paragraph in document.paragraphs:
            add(paragraph.text)
            if truncated:
                break
        if not truncated:
            for table in document.tables:
                for row in table.rows:
                    add(" | ".join(cell.text for cell in row.cells))
                    if truncated:
                        break
                if truncated:
                    break
        props = document.core_properties
        return ExtractedContent("\n".join(parts), "DOCX", props.title or "", props.author or "",
                                props.subject or "", truncated=truncated)

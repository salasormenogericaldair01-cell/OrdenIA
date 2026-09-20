"""Small, generated documents exercise each local extractor."""

from pathlib import Path
from zipfile import ZipFile

import pymupdf
import pytest
from docx import Document
from openpyxl import Workbook
from pptx import Presentation

from ordenia.analysis import AnalysisLimits, ContentAnalyzer
from ordenia.analysis.registry import CODE_EXTENSIONS, TEXT_EXTENSIONS


def test_pdf_text_metadata_and_image_only(tmp_path: Path) -> None:
    path = tmp_path / "manual.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((70, 70), "ESP32 sensor ultrasonico")
        document.set_metadata({"title": "Manual de placa", "author": "OrdenIA", "subject": "Sensores"})
        document.save(path)
    result = ContentAnalyzer().analyze(path)
    assert result.status == "indexed"
    assert result.page_count == 1
    assert result.title == "Manual de placa"
    assert result.author == "OrdenIA"
    assert result.subject == "Sensores"
    assert "ESP32" in result.text
    assert "ESP32" in result.keywords

    blank = tmp_path / "imagen.pdf"
    with pymupdf.open() as document:
        document.new_page()
        document.save(blank)
    result = ContentAnalyzer().analyze(blank)
    assert result.status == "no_text"
    assert result.error == "PDF sin texto extraíble"


def test_pdf_corrupt_and_encrypted_are_failed_without_crashing(tmp_path: Path) -> None:
    broken = tmp_path / "broken.pdf"
    broken.write_bytes(b"not a pdf")
    assert ContentAnalyzer().analyze(broken).status == "failed"
    protected = tmp_path / "protected.pdf"
    with pymupdf.open() as document:
        document.new_page().insert_text((70, 70), "secreto")
        document.save(protected, encryption=pymupdf.PDF_ENCRYPT_AES_256, owner_pw="owner", user_pw="secret")
    result = ContentAnalyzer().analyze(protected)
    assert result.status == "failed"
    assert "cifrado" in result.error


def test_docx_paragraphs_header_tables_and_properties(tmp_path: Path) -> None:
    path = tmp_path / "informe.docx"
    document = Document()
    document.add_heading("Inventario institucional", 1)
    document.add_paragraph("Sensor de residuos disponible")
    document.sections[0].header.paragraphs[0].text = "Encabezado del proyecto"
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Código"
    table.cell(0, 1).text = "ESP32"
    document.core_properties.title = "Informe de inventario"
    document.save(path)
    result = ContentAnalyzer().analyze(path)
    assert result.status == "indexed"
    assert all(term in result.text for term in ("Encabezado", "Inventario institucional", "Sensor", "ESP32"))
    assert result.title == "Informe de inventario"


def test_xlsx_streams_sheet_names_values_without_evaluating_formulas(tmp_path: Path) -> None:
    path = tmp_path / "datos.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Inventario"
    sheet.append(["Componente", "Cantidad"])
    sheet.append(["sensor ultrasonico", 3])
    sheet.append(["=1+2", None])
    workbook.save(path)
    result = ContentAnalyzer().analyze(path)
    assert result.status == "indexed"
    assert "Inventario" in result.text
    assert "sensor ultrasonico" in result.text
    assert "Cantidad" in result.text
    assert "3" in result.text
    assert "=1+2" not in result.text


def test_pptx_extracts_slide_text_and_count(tmp_path: Path) -> None:
    path = tmp_path / "presentacion.pptx"
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Proyecto residuos"
    slide.placeholders[1].text = "Recolección inteligente con ESP32"
    presentation.core_properties.title = "Presentación de proyecto"
    presentation.save(path)
    result = ContentAnalyzer().analyze(path)
    assert result.status == "indexed"
    assert result.slide_count == 1
    assert "ESP32" in result.text
    assert result.title == "Presentación de proyecto"


@pytest.mark.parametrize("suffix", [".txt", ".md", ".log", ".csv", ".py", ".ino", ".json", ".toml", ".ps1"])
def test_text_markdown_and_code_are_read_as_data(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / ("notas" + suffix)
    path.write_text("ESP32 sensor ultrasónico\nraise RuntimeError('never execute')", encoding="utf-8")
    result = ContentAnalyzer().analyze(path)
    assert result.status == "indexed"
    assert "ESP32" in result.text
    assert "never execute" in result.text
    assert result.extractor == ("Texto" if suffix in TEXT_EXTENSIONS else "Código")


def test_cp1252_fallback_and_limits(tmp_path: Path) -> None:
    path = tmp_path / "notas.txt"
    path.write_bytes("recolección local".encode("cp1252"))
    assert "recolección" in ContentAnalyzer().analyze(path).text
    short = ContentAnalyzer(AnalysisLimits(max_extracted_characters=5)).analyze(path)
    assert short.truncated and len(short.text) == 5
    too_large = ContentAnalyzer(AnalysisLimits(max_file_bytes=2)).analyze(path)
    assert too_large.status == "skipped"
    assert "demasiado grande" in too_large.error


@pytest.mark.parametrize("suffix", [".doc", ".xls", ".ppt"])
def test_legacy_formats_are_explicitly_unsupported(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / ("legacy" + suffix)
    path.write_text("placeholder")
    result = ContentAnalyzer().analyze(path)
    assert result.status == "unsupported"
    assert result.error == "No compatible con análisis de contenido en V0.3"


def test_registry_covers_requested_safe_source_extensions() -> None:
    requested = {".py", ".js", ".ts", ".tsx", ".jsx", ".java", ".c", ".cpp", ".h", ".hpp",
                 ".ino", ".html", ".css", ".json", ".yaml", ".yml", ".toml", ".sql", ".sh",
                 ".ps1", ".bat", ".xml"}
    assert requested <= CODE_EXTENSIONS


def test_office_archive_expansion_is_bounded(tmp_path: Path) -> None:
    path = tmp_path / "compressed.docx"
    with ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", "A" * 2000)
    result = ContentAnalyzer(AnalysisLimits(max_zip_uncompressed_bytes=1000)).analyze(path)
    assert result.status == "skipped"
    assert "comprimido demasiado grande" in result.error
